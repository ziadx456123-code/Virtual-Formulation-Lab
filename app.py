from __future__ import annotations

import io
import json
import warnings
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Any

import joblib
import numpy as np
import pandas as pd
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go

from scipy.optimize import differential_evolution

from sklearn.ensemble import (
    ExtraTreesRegressor,
    RandomForestRegressor,
    GradientBoostingRegressor,
    HistGradientBoostingRegressor,
)
from sklearn.feature_selection import SelectKBest, mutual_info_regression
from sklearn.impute import SimpleImputer
from sklearn.inspection import permutation_importance
from sklearn.linear_model import Ridge
from sklearn.metrics import r2_score, mean_absolute_error, mean_squared_error
from sklearn.model_selection import (
    train_test_split,
    RepeatedKFold,
    RandomizedSearchCV,
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVR

warnings.filterwarnings("ignore")


# ============================================================
# CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Virtual Formulation Lab",
    page_icon="🧪",
    layout="wide",
    initial_sidebar_state="collapsed",
)

RANDOM_SEED = 42
TEST_SIZE = 0.20
CV_FOLDS = 5
CV_REPEATS = 2
N_ITER_SEARCH = 20
MAX_OPT_VARS = 12

BASE_DIR = Path(__file__).resolve().parent

MODEL_DIR = BASE_DIR / "saved_cqa_models"
MODEL_DIR.mkdir(exist_ok=True)

DATASET_CANDIDATES = [
    BASE_DIR / "final Data All Exipients.csv",
    BASE_DIR / "final Data All Exipients.xlsx",
    BASE_DIR / "final Data All Exipients.xls",
]

TARGETS = [
    "HARDNESS",
    "FRIABILITY",
    "Drug content",
    "Water absorption ratio",
    "DISINTEGRATION_TIME",
]


# ============================================================
# PROFESSIONAL UI STYLE
# ============================================================

st.markdown(
    """
    <style>

    #MainMenu {
        visibility: hidden;
    }

    footer {
        visibility: hidden;
    }

    header {
        visibility: hidden;
    }

    .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
        max-width: 1400px;
    }

    .main-title {
        font-size: 2.4rem;
        font-weight: 700;
        margin-bottom: 0.2rem;
    }

    .subtitle {
        color: #6b7280;
        font-size: 1rem;
        margin-bottom: 1.8rem;
    }

    .section-title {
        font-size: 1.35rem;
        font-weight: 650;
        margin-top: 1rem;
        margin-bottom: 0.7rem;
    }

    .result-card {
        border: 1px solid #e5e7eb;
        border-radius: 14px;
        padding: 18px;
        background: #ffffff;
        box-shadow: 0 2px 8px rgba(0,0,0,0.04);
        min-height: 120px;
    }

    .result-label {
        color: #6b7280;
        font-size: 0.88rem;
        margin-bottom: 8px;
    }

    .result-value {
        font-size: 1.45rem;
        font-weight: 700;
    }

    .info-box {
        padding: 14px 16px;
        border-radius: 12px;
        background: #f8fafc;
        border: 1px solid #e2e8f0;
        color: #475569;
        margin: 12px 0;
    }

    .success-box {
        padding: 14px 16px;
        border-radius: 12px;
        background: #f0fdf4;
        border: 1px solid #bbf7d0;
        color: #166534;
        margin: 12px 0;
    }

    .warning-box {
        padding: 14px 16px;
        border-radius: 12px;
        background: #fffbeb;
        border: 1px solid #fde68a;
        color: #92400e;
        margin: 12px 0;
    }

    div[data-testid="stMetric"] {
        border: 1px solid #e5e7eb;
        padding: 14px;
        border-radius: 12px;
        background: #ffffff;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# DATA HANDLING
# ============================================================

@st.cache_data(show_spinner=False)
def load_dataset() -> pd.DataFrame:

    for path in DATASET_CANDIDATES:

        if path.exists():

            if path.suffix.lower() == ".csv":
                df = pd.read_csv(path)

            else:
                df = pd.read_excel(path)

            return clean_dataset(df)

    raise FileNotFoundError(
        "Required formulation dataset was not found."
    )


def clean_dataset(df: pd.DataFrame) -> pd.DataFrame:

    df = df.copy()

    # Convert numeric-looking object columns
    for col in df.columns:

        if df[col].dtype == "object":

            converted = pd.to_numeric(
                df[col].astype(str).str.replace(",", "", regex=False),
                errors="coerce",
            )

            valid_ratio = converted.notna().mean()

            if valid_ratio >= 0.80:
                df[col] = converted

    missing_targets = [
        c for c in TARGETS if c not in df.columns
    ]

    if missing_targets:
        raise ValueError(
            f"Required CQA columns are missing: {missing_targets}"
        )

    numeric_cols = df.select_dtypes(
        include=[np.number]
    ).columns.tolist()

    predictors = [
        c for c in numeric_cols
        if c not in TARGETS
    ]

    if not predictors:
        raise ValueError("No numeric formulation inputs were found.")

    df = df[predictors + TARGETS]

    df = df.replace(
        [np.inf, -np.inf],
        np.nan
    )

    df = df.drop_duplicates()

    # Remove rows where all targets are missing
    df = df.dropna(
        subset=TARGETS,
        how="all"
    )

    return df


# ============================================================
# METRICS
# ============================================================

def calculate_metrics(
    y_true,
    y_pred
) -> Dict[str, float]:

    rmse = np.sqrt(
        mean_squared_error(y_true, y_pred)
    )

    return {
        "R2": float(r2_score(y_true, y_pred)),
        "MAE": float(mean_absolute_error(y_true, y_pred)),
        "RMSE": float(rmse),
    }


# ============================================================
# MODEL DEFINITIONS
# ============================================================

def get_model_candidates():

    candidates = {

        "Extra Trees": (
            ExtraTreesRegressor(
                random_state=RANDOM_SEED,
                n_jobs=-1
            ),
            {
                "model__n_estimators": [300, 500, 700, 900],
                "model__max_depth": [
                    None, 5, 8, 12, 16, 20
                ],
                "model__min_samples_split": [
                    2, 3, 4, 5
                ],
                "model__min_samples_leaf": [
                    1, 2, 3
                ],
                "model__max_features": [
                    0.5, 0.7, 0.9, 1.0
                ],
            },
        ),

        "Random Forest": (
            RandomForestRegressor(
                random_state=RANDOM_SEED,
                n_jobs=-1
            ),
            {
                "model__n_estimators": [300, 500, 700, 900],
                "model__max_depth": [
                    None, 5, 8, 12, 16, 20
                ],
                "model__min_samples_split": [
                    2, 3, 4, 5
                ],
                "model__min_samples_leaf": [
                    1, 2, 3
                ],
                "model__max_features": [
                    0.5, 0.7, 0.9, 1.0
                ],
            },
        ),

        "Gradient Boosting": (
            GradientBoostingRegressor(
                random_state=RANDOM_SEED
            ),
            {
                "model__n_estimators": [
                    100, 200, 300, 500
                ],
                "model__learning_rate": [
                    0.01, 0.03, 0.05, 0.08, 0.1
                ],
                "model__max_depth": [
                    1, 2, 3, 4
                ],
                "model__min_samples_split": [
                    2, 3, 4
                ],
                "model__min_samples_leaf": [
                    1, 2, 3
                ],
            },
        ),

        "HistGradient Boosting": (
            HistGradientBoostingRegressor(
                random_state=RANDOM_SEED
            ),
            {
                "model__max_iter": [
                    100, 200, 300, 500
                ],
                "model__learning_rate": [
                    0.02, 0.05, 0.08, 0.1
                ],
                "model__max_leaf_nodes": [
                    7, 15, 31
                ],
                "model__l2_regularization": [
                    0.0, 0.1, 0.5, 1.0
                ],
            },
        ),

        "SVR": (
            SVR(kernel="rbf"),
            {
                "model__C": [
                    0.1, 1, 10, 30, 100, 300
                ],
                "model__gamma": [
                    "scale",
                    "auto",
                    0.001,
                    0.01,
                    0.1
                ],
                "model__epsilon": [
                    0.001,
                    0.01,
                    0.05,
                    0.1,
                    0.2
                ],
            },
        ),

        "Ridge": (
            Ridge(),
            {
                "model__alpha": [
                    0.01,
                    0.1,
                    1,
                    10,
                    50,
                    100
                ],
            },
        ),
    }

    return candidates


# ============================================================
# PIPELINE
# ============================================================

def make_pipeline(model):

    return Pipeline(
        steps=[
            (
                "imputer",
                SimpleImputer(
                    strategy="median",
                    add_indicator=True
                )
            ),

            (
                "feature_selection",
                SelectKBest(
                    score_func=mutual_info_regression,
                    k="all"
                )
            ),

            (
                "scaler",
                StandardScaler()
            ),

            (
                "model",
                model
            ),
        ]
    )


# ============================================================
# MODEL TRAINING
# ============================================================

def train_one_cqa(
    df: pd.DataFrame,
    target: str
):

    X = df.drop(
        columns=TARGETS
    )

    y = df[target]

    valid = y.notna()

    X = X.loc[valid].copy()
    y = y.loc[valid].copy()

    if len(X) < 30:
        raise ValueError(
            f"Not enough observations for {target}."
        )

    X_train, X_test, y_train, y_test = train_test_split(
        X,
        y,
        test_size=TEST_SIZE,
        random_state=RANDOM_SEED
    )

    cv = RepeatedKFold(
        n_splits=CV_FOLDS,
        n_repeats=CV_REPEATS,
        random_state=RANDOM_SEED
    )

    candidates = get_model_candidates()

    leaderboard = []

    best_pipeline = None
    best_model_name = None
    best_cv_score = -np.inf

    for model_name, (model, params) in candidates.items():

        pipeline = make_pipeline(model)

        search = RandomizedSearchCV(
            estimator=pipeline,
            param_distributions=params,
            n_iter=N_ITER_SEARCH,
            scoring="r2",
            cv=cv,
            random_state=RANDOM_SEED,
            n_jobs=-1,
            refit=True,
            error_score=np.nan,
        )

        try:

            search.fit(X_train, y_train)

            score = search.best_score_

            if not np.isfinite(score):
                continue

            leaderboard.append(
                {
                    "Model": model_name,
                    "CV R2": float(score),
                }
            )

            if score > best_cv_score:

                best_cv_score = score
                best_pipeline = search.best_estimator_
                best_model_name = model_name

        except Exception:
            continue

    if best_pipeline is None:
        raise RuntimeError(
            f"No model could be trained successfully for {target}."
        )

    y_pred = best_pipeline.predict(X_test)

    metrics = calculate_metrics(
        y_test,
        y_pred
    )

    metrics["CV R2"] = float(best_cv_score)

    # --------------------------------------------------------
    # Scientific interpretability
    # --------------------------------------------------------

    try:

        perm = permutation_importance(
            best_pipeline,
            X_train,
            y_train,
            scoring="r2",
            n_repeats=15,
            random_state=RANDOM_SEED,
            n_jobs=-1,
        )

        importance = pd.DataFrame(
            {
                "Feature": X_train.columns,
                "Importance": np.maximum(
                    perm.importances_mean,
                    0
                ),
            }
        )

        total = importance["Importance"].sum()

        if total > 0:

            importance["Importance_%"] = (
                importance["Importance"] / total * 100
            )

        else:

            importance["Importance_%"] = 0.0

        importance = importance.sort_values(
            "Importance_%",
            ascending=False
        )

    except Exception:

        importance = pd.DataFrame(
            columns=[
                "Feature",
                "Importance",
                "Importance_%"
            ]
        )

    artifact = {
        "model": best_pipeline,
        "target": target,
        "features": X.columns.tolist(),
        "metrics": metrics,
        "best_model": best_model_name,
        "importance": importance,
        "leaderboard": pd.DataFrame(leaderboard)
        .sort_values("CV R2", ascending=False),
        "X_test": X_test,
        "y_test": y_test,
        "y_pred": y_pred,
    }

    return artifact


# ============================================================
# MODEL STORAGE
# ============================================================

def model_path(target: str):

    safe_name = (
        target
        .replace(" ", "_")
        .replace("/", "_")
    )

    return MODEL_DIR / f"{safe_name}.joblib"


def save_artifact(
    target: str,
    artifact: Dict[str, Any]
):

    joblib.dump(
        artifact,
        model_path(target)
    )


def load_artifact(
    target: str
):

    path = model_path(target)

    if not path.exists():
        return None

    try:
        return joblib.load(path)
    except Exception:
        return None


def all_models_available():

    return all(
        model_path(target).exists()
        for target in TARGETS
    )


# ============================================================
# AUTOMATIC AI INITIALIZATION
# ============================================================

def initialize_models():

    df = load_dataset()

    progress = st.progress(0)

    status = st.empty()

    total = len(TARGETS)

    for i, target in enumerate(TARGETS):

        status.info(
            f"Preparing prediction engine for {target}..."
        )

        artifact = train_one_cqa(
            df,
            target
        )

        save_artifact(
            target,
            artifact
        )

        progress.progress(
            (i + 1) / total
        )

    status.success(
        "Prediction engine is ready."
    )

    progress.empty()


@st.cache_resource(show_spinner=False)
def load_all_models():

    models = {}

    for target in TARGETS:

        artifact = load_artifact(target)

        if artifact is not None:
            models[target] = artifact

    return models


def ensure_models():

    if not all_models_available():

        with st.spinner(
            "Initializing the formulation prediction engine..."
        ):
            initialize_models()

        st.cache_resource.clear()

        return load_all_models()

    return load_all_models()


# ============================================================
# PREDICTION
# ============================================================

def predict_cqas(
    models: Dict[str, Any],
    inputs: pd.DataFrame
):

    predictions = {}

    for target in TARGETS:

        artifact = models[target]

        model = artifact["model"]

        predictions[target] = float(
            model.predict(inputs)[0]
        )

    return predictions


# ============================================================
# DOMAIN INFORMATION
# ============================================================

def get_input_ranges(
    df: pd.DataFrame
):

    ranges = {}

    for col in df.columns:

        if col in TARGETS:
            continue

        series = pd.to_numeric(
            df[col],
            errors="coerce"
        ).dropna()

        if len(series) == 0:
            continue

        ranges[col] = {
            "min": float(series.min()),
            "max": float(series.max()),
            "unique": series.nunique(),
            "values": sorted(
                series.unique().tolist()
            )
        }

    return ranges


# ============================================================
# DESIRABILITY
# ============================================================

def desirability(
    value: float,
    goal: str,
    low: float,
    high: float,
    target: float | None = None
):

    if not np.isfinite(value):
        return 0.0

    if high <= low:
        return 1.0

    if goal == "Minimize":

        if value <= low:
            return 1.0

        if value >= high:
            return 0.0

        return (high - value) / (high - low)

    if goal == "Maximize":

        if value >= high:
            return 1.0

        if value <= low:
            return 0.0

        return (value - low) / (high - low)

    # Target
    if target is None:
        target = (low + high) / 2

    target = float(
        np.clip(target, low, high)
    )

    if value == target:
        return 1.0

    if value < target:

        if target == low:
            return 1.0

        return max(
            0.0,
            (value - low) / (target - low)
        )

    if high == target:
        return 1.0

    return max(
        0.0,
        (high - value) / (high - target)
    )


def overall_desirability(
    predictions: Dict[str, float],
    goals: Dict[str, Dict[str, float]],
    df: pd.DataFrame
):

    values = []

    for target in TARGETS:

        y = df[target].dropna()

        low = float(y.min())
        high = float(y.max())

        config = goals[target]

        d = desirability(
            predictions[target],
            config["goal"],
            low,
            high,
            config.get("target")
        )

        values.append(
            np.clip(d, 0, 1)
        )

    if not values:
        return 0.0

    return float(
        np.prod(values) ** (1 / len(values))
    )


# ============================================================
# OPTIMIZATION VARIABLES
# ============================================================

def select_optimization_variables(
    models
):

    score_map = {}

    for target in TARGETS:

        importance = models[target]["importance"]

        for _, row in importance.iterrows():

            feature = row["Feature"]

            score_map[feature] = (
                score_map.get(feature, 0)
                + float(row["Importance_%"])
            )

    ranked = sorted(
        score_map.items(),
        key=lambda x: x[1],
        reverse=True
    )

    return [
        feature
        for feature, _ in ranked[:MAX_OPT_VARS]
    ]


# ============================================================
# OPTIMIZATION
# ============================================================

def optimize_formulation(
    df,
    models,
    baseline_inputs,
    goals
):

    ranges = get_input_ranges(df)

    selected = select_optimization_variables(
        models
    )

    selected = [
        x for x in selected
        if x in ranges
    ]

    if not selected:
        return baseline_inputs.copy(), {}, 0.0

    base = baseline_inputs.copy()

    bounds = [
        (
            ranges[f]["min"],
            ranges[f]["max"]
        )
        for f in selected
    ]

    discrete_values = {}

    for feature in selected:

        if ranges[feature]["unique"] <= 10:

            discrete_values[feature] = (
                ranges[feature]["values"]
            )

    def create_candidate(values):

        candidate = base.copy()

        for feature, value in zip(
            selected,
            values
        ):

            if feature in discrete_values:

                observed = np.array(
                    discrete_values[feature]
                )

                value = float(
                    observed[
                        np.argmin(
                            np.abs(observed - value)
                        )
                    ]
                )

            candidate[feature] = value

        return candidate

    def objective(values):

        candidate = create_candidate(
            values
        )

        predictions = predict_cqas(
            models,
            candidate
        )

        d = overall_desirability(
            predictions,
            goals,
            df
        )

        return -d

    result = differential_evolution(
        objective,
        bounds=bounds,
        seed=RANDOM_SEED,
        maxiter=80,
        popsize=10,
        polish=True,
        updating="immediate",
        workers=1,
    )

    optimized_inputs = create_candidate(
        result.x
    )

    optimized_predictions = predict_cqas(
        models,
        optimized_inputs
    )

    optimized_desirability = overall_desirability(
        optimized_predictions,
        goals,
        df
    )

    return (
        optimized_inputs,
        optimized_predictions,
        optimized_desirability
    )


# ============================================================
# EXCEL REPORT
# ============================================================

def format_excel(writer):

    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    workbook = writer.book

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        for column_cells in worksheet.columns:

            max_length = 0

            for cell in column_cells:

                try:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )
                except Exception:
                    pass

            width = min(
                max(max_length + 2, 10),
                40
            )

            worksheet.column_dimensions[
                get_column_letter(
                    column_cells[0].column
                )
            ].width = width


def create_excel_report(
    formulation_name,
    predictions,
    optimized_predictions,
    baseline_inputs,
    optimized_inputs,
    models,
    optimization_desirability,
    goals
):

    output = io.BytesIO()

    timestamp = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        summary = pd.DataFrame(
            {
                "Item": [
                    "Formulation Name",
                    "Report Date & Time",
                    "Overall Optimization Desirability",
                ],
                "Value": [
                    formulation_name,
                    timestamp,
                    optimization_desirability,
                ],
            }
        )

        summary.to_excel(
            writer,
            index=False,
            sheet_name="Report Summary"
        )

        prediction_rows = []

        for target in TARGETS:

            prediction_rows.append(
                {
                    "CQA": target,
                    "Baseline Prediction":
                        predictions.get(target),
                    "Optimized Prediction":
                        optimized_predictions.get(target),
                    "Goal":
                        goals[target]["goal"],
                    "Target":
                        goals[target].get("target"),
                }
            )

        pd.DataFrame(
            prediction_rows
        ).to_excel(
            writer,
            index=False,
            sheet_name="CQA Results"
        )

        inputs_rows = []

        for feature in baseline_inputs.columns:

            inputs_rows.append(
                {
                    "Input": feature,
                    "Baseline":
                        baseline_inputs.iloc[0][feature],
                    "Optimized":
                        optimized_inputs.iloc[0][feature],
                }
            )

        pd.DataFrame(
            inputs_rows
        ).to_excel(
            writer,
            index=False,
            sheet_name="Formulation Inputs"
        )

        validation_rows = []

        for target in TARGETS:

            metrics = models[target]["metrics"]

            validation_rows.append(
                {
                    "CQA": target,
                    "CV R2":
                        metrics.get("CV R2"),
                    "Test R2":
                        metrics.get("R2"),
                    "Test RMSE":
                        metrics.get("RMSE"),
                    "Test MAE":
                        metrics.get("MAE"),
                }
            )

        pd.DataFrame(
            validation_rows
        ).to_excel(
            writer,
            index=False,
            sheet_name="Scientific Validation"
        )

        importance_rows = []

        for target in TARGETS:

            imp = models[target]["importance"].copy()

            imp.insert(
                0,
                "CQA",
                target
            )

            importance_rows.append(
                imp
            )

        if importance_rows:

            pd.concat(
                importance_rows,
                ignore_index=True
            ).to_excel(
                writer,
                index=False,
                sheet_name="Feature Importance"
            )

        format_excel(writer)

    output.seek(0)

    return output


# ============================================================
# SESSION STATE
# ============================================================

if "formulation_name" not in st.session_state:
    st.session_state.formulation_name = "Formulation 01"

if "predictions" not in st.session_state:
    st.session_state.predictions = None

if "baseline_inputs" not in st.session_state:
    st.session_state.baseline_inputs = None

if "optimized_inputs" not in st.session_state:
    st.session_state.optimized_inputs = None

if "optimized_predictions" not in st.session_state:
    st.session_state.optimized_predictions = None


# ============================================================
# LOAD DATA / MODELS
# ============================================================

try:

    df = load_dataset()

except Exception as e:

    st.error(
        "The formulation data could not be loaded."
    )

    st.stop()


try:

    models = ensure_models()

except Exception as e:

    st.error(
        "The prediction engine could not be initialized."
    )

    st.stop()


if not all(
    target in models
    for target in TARGETS
):

    st.error(
        "The prediction engine is incomplete."
    )

    st.stop()


input_columns = [
    c for c in df.columns
    if c not in TARGETS
]

input_ranges = get_input_ranges(df)


# ============================================================
# HEADER
# ============================================================

st.markdown(
    '<div class="main-title">🧪 Virtual Formulation Lab</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">'
    'AI-assisted pharmaceutical formulation prediction and optimization'
    '</div>',
    unsafe_allow_html=True
)


# ============================================================
# MAIN NAVIGATION
# ============================================================

tab_lab, tab_optimizer, tab_insights = st.tabs(
    [
        "🧪 Virtual Formulation Lab",
        "🎯 AI Optimizer",
        "📊 Scientific Insights",
    ]
)


# ============================================================
# TAB 1 — VIRTUAL FORMULATION LAB
# ============================================================

with tab_lab:

    st.markdown(
        '<div class="section-title">'
        'Formulation Setup'
        '</div>',
        unsafe_allow_html=True
    )

    formulation_name = st.text_input(
        "Formulation Name",
        value=st.session_state.formulation_name,
        key="formulation_name_input"
    )

    st.session_state.formulation_name = (
        formulation_name
    )

    st.markdown(
        '<div class="info-box">'
        'Enter formulation and process parameters within the '
        'experimental ranges represented by the model.'
        '</div>',
        unsafe_allow_html=True
    )

    values = {}

    # Organize inputs in columns
    for start in range(0, len(input_columns), 3):

        cols = st.columns(3)

        for col_index, feature in enumerate(
            input_columns[start:start + 3]
        ):

            info = input_ranges[feature]

            minimum = info["min"]
            maximum = info["max"]

            if minimum == maximum:
                values[feature] = minimum

                with cols[col_index]:
                    st.number_input(
                        feature,
                        value=float(minimum),
                        disabled=True,
                        key=f"input_{feature}"
                    )

                continue

            unique_values = info["values"]

            with cols[col_index]:

                # Categorical-like numeric input
                if (
                    info["unique"] <= 10
                    and len(unique_values) <= 10
                ):

                    selected = st.selectbox(
                        feature,
                        options=unique_values,
                        key=f"input_{feature}"
                    )

                    values[feature] = selected

                else:

                    values[feature] = st.number_input(
                        feature,
                        min_value=float(minimum),
                        max_value=float(maximum),
                        value=float(
                            (minimum + maximum) / 2
                        ),
                        key=f"input_{feature}"
                    )

    st.markdown("")

    predict_clicked = st.button(
        "🔬 Predict CQAs",
        type="primary",
        use_container_width=True
    )

    if predict_clicked:

        baseline_inputs = pd.DataFrame(
            [values],
            columns=input_columns
        )

        predictions = predict_cqas(
            models,
            baseline_inputs
        )

        st.session_state.baseline_inputs = (
            baseline_inputs
        )

        st.session_state.predictions = (
            predictions
        )

    if st.session_state.predictions:

        st.markdown(
            '<div class="section-title">'
            'Predicted Critical Quality Attributes'
            '</div>',
            unsafe_allow_html=True
        )

        predictions = (
            st.session_state.predictions
        )

        result_cols = st.columns(5)

        for i, target in enumerate(TARGETS):

            value = predictions[target]

            with result_cols[i]:

                st.markdown(
                    f"""
                    <div class="result-card">
                        <div class="result-label">
                            {target}
                        </div>
                        <div class="result-value">
                            {value:.4g}
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True
                )

        st.markdown(
            """
            <div class="warning-box">
            <b>Important:</b> These are model-based predictions,
            not laboratory measurements. Experimental confirmation
            is required before making formulation decisions.
            </div>
            """,
            unsafe_allow_html=True
        )

        # Export current prediction
        prediction_export = pd.DataFrame(
            [
                {
                    "CQA": target,
                    "Predicted Value":
                        predictions[target]
                }
                for target in TARGETS
            ]
        )

        st.download_button(
            "📥 Export Prediction",
            data=prediction_export.to_csv(
                index=False
            ).encode("utf-8"),
            file_name=(
                f"{formulation_name}_prediction.csv"
            ),
            mime="text/csv",
        )


# ============================================================
# TAB 2 — AI OPTIMIZER
# ============================================================

with tab_optimizer:

    st.markdown(
        '<div class="section-title">'
        'Define CQA Objectives'
        '</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="info-box">'
        'Set the desired direction for each CQA. '
        'For Target objectives, enter the desired target value.'
        '</div>',
        unsafe_allow_html=True
    )

    goals = {}

    goal_columns = st.columns(2)

    for i, target in enumerate(TARGETS):

        with goal_columns[i % 2]:

            goal = st.selectbox(
                f"{target} objective",
                [
                    "Target",
                    "Minimize",
                    "Maximize",
                ],
                key=f"goal_{target}"
            )

            target_value = None

            if goal == "Target":

                series = df[target].dropna()

                target_value = st.number_input(
                    f"{target} target",
                    min_value=float(series.min()),
                    max_value=float(series.max()),
                    value=float(series.median()),
                    key=f"target_{target}"
                )

            goals[target] = {
                "goal": goal,
                "target": target_value,
            }

    st.markdown("")

    optimize_clicked = st.button(
        "🚀 Find Optimal Formulation",
        type="primary",
        use_container_width=True
    )

    if optimize_clicked:

        if (
            st.session_state.baseline_inputs
            is None
        ):

            st.warning(
                "Please create a formulation prediction first "
                "from the Virtual Formulation Lab."
            )

        else:

            with st.spinner(
                "Searching for the best formulation within the experimental domain..."
            ):

                optimized_inputs, optimized_predictions, d = (
                    optimize_formulation(
                        df,
                        models,
                        st.session_state.baseline_inputs,
                        goals,
                    )
                )

            st.session_state.optimized_inputs = (
                optimized_inputs
            )

            st.session_state.optimized_predictions = (
                optimized_predictions
            )

            st.session_state.optimized_desirability = d

    if (
        st.session_state.optimized_inputs is not None
        and st.session_state.optimized_predictions is not None
    ):

        optimized_inputs = (
            st.session_state.optimized_inputs
        )

        optimized_predictions = (
            st.session_state.optimized_predictions
        )

        baseline_inputs = (
            st.session_state.baseline_inputs
        )

        st.markdown(
            '<div class="section-title">'
            'AI Recommended Formulation'
            '</div>',
            unsafe_allow_html=True
        )

        comparison_inputs = pd.DataFrame(
            {
                "Input": input_columns,
                "Current": [
                    baseline_inputs.iloc[0][c]
                    for c in input_columns
                ],
                "Recommended": [
                    optimized_inputs.iloc[0][c]
                    for c in input_columns
                ],
            }
        )

        st.dataframe(
            comparison_inputs,
            use_container_width=True,
            hide_index=True
        )

        st.markdown(
            '<div class="section-title">'
            'CQA Prediction Comparison'
            '</div>',
            unsafe_allow_html=True
        )

        comparison_cqa = pd.DataFrame(
            {
                "CQA": TARGETS,
                "Current": [
                    st.session_state.predictions[t]
                    for t in TARGETS
                ],
                "Optimized": [
                    optimized_predictions[t]
                    for t in TARGETS
                ],
            }
        )

        st.dataframe(
            comparison_cqa,
            use_container_width=True,
            hide_index=True
        )

        d = st.session_state.get(
            "optimized_desirability",
            0.0
        )

        st.metric(
            "Overall Desirability",
            f"{d:.3f}"
        )

        # Radar
        baseline_d = []

        optimized_d = []

        for target in TARGETS:

            series = df[target].dropna()

            low = float(series.min())
            high = float(series.max())

            config = goals[target]

            baseline_d.append(
                desirability(
                    st.session_state.predictions[target],
                    config["goal"],
                    low,
                    high,
                    config.get("target")
                )
            )

            optimized_d.append(
                desirability(
                    optimized_predictions[target],
                    config["goal"],
                    low,
                    high,
                    config.get("target")
                )
            )

        fig = go.Figure()

        fig.add_trace(
            go.Scatterpolar(
                r=baseline_d + [baseline_d[0]],
                theta=TARGETS + [TARGETS[0]],
                fill="toself",
                name="Current"
            )
        )

        fig.add_trace(
            go.Scatterpolar(
                r=optimized_d + [optimized_d[0]],
                theta=TARGETS + [TARGETS[0]],
                fill="toself",
                name="Optimized"
            )
        )

        fig.update_layout(
            polar=dict(
                radialaxis=dict(
                    visible=True,
                    range=[0, 1]
                )
            ),
            showlegend=True,
            margin=dict(
                l=40,
                r=40,
                t=40,
                b=40
            ),
        )

        st.plotly_chart(
            fig,
            use_container_width=True
        )

        st.markdown(
            """
            <div class="warning-box">
            <b>Scientific note:</b> The recommendation is generated
            within the historical experimental domain represented by
            the available data. It is not a validated Design Space
            and requires laboratory confirmation.
            </div>
            """,
            unsafe_allow_html=True
        )

        # Full Excel report
        report = create_excel_report(
            formulation_name,
            st.session_state.predictions,
            optimized_predictions,
            baseline_inputs,
            optimized_inputs,
            models,
            d,
            goals,
        )

        st.download_button(
            "📥 Export Full Scientific Report",
            data=report,
            file_name=(
                f"{formulation_name}_QbD_Report.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True
        )


# ============================================================
# TAB 3 — SCIENTIFIC INSIGHTS
# ============================================================

with tab_insights:

    st.markdown(
        '<div class="section-title">'
        'Scientific Model Performance'
        '</div>',
        unsafe_allow_html=True
    )

    validation_rows = []

    for target in TARGETS:

        metrics = models[target]["metrics"]

        validation_rows.append(
            {
                "CQA": target,
                "CV R²": metrics["CV R2"],
                "Test R²": metrics["R2"],
                "Test RMSE": metrics["RMSE"],
            }
        )

    validation_df = pd.DataFrame(
        validation_rows
    )

    st.dataframe(
        validation_df.style.format(
            {
                "CV R²": "{:.3f}",
                "Test R²": "{:.3f}",
                "Test RMSE": "{:.4g}",
            }
        ),
        use_container_width=True,
        hide_index=True
    )

    st.markdown(
        '<div class="info-box">'
        '<b>Interpretation:</b> R² indicates how well the model '
        'explains variation in the measured CQA. RMSE represents '
        'prediction error in the same numerical scale as the CQA. '
        'These metrics describe model performance and are not '
        'classification accuracy percentages.'
        '</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="section-title">'
        'Predicted vs Actual'
        '</div>',
        unsafe_allow_html=True
    )

    selected_cqa = st.selectbox(
        "Select CQA",
        TARGETS,
        key="validation_cqa"
    )

    artifact = models[selected_cqa]

    y_test = artifact["y_test"]
    y_pred = artifact["y_pred"]

    scatter_df = pd.DataFrame(
        {
            "Actual": y_test.values,
            "Predicted": y_pred,
        }
    )

    fig = px.scatter(
        scatter_df,
        x="Actual",
        y="Predicted",
        title=selected_cqa,
    )

    minimum = min(
        scatter_df["Actual"].min(),
        scatter_df["Predicted"].min()
    )

    maximum = max(
        scatter_df["Actual"].max(),
        scatter_df["Predicted"].max()
    )

    fig.add_shape(
        type="line",
        x0=minimum,
        y0=minimum,
        x1=maximum,
        y1=maximum,
        line=dict(
            dash="dash"
        )
    )

    fig.update_layout(
        margin=dict(
            l=40,
            r=40,
            t=60,
            b=40
        )
    )

    st.plotly_chart(
        fig,
        use_container_width=True
    )

    st.markdown(
        '<div class="section-title">'
        'Key Formulation Drivers'
        '</div>',
        unsafe_allow_html=True
    )

    importance = artifact[
        "importance"
    ].head(15).copy()

    if not importance.empty:

        importance = importance.sort_values(
            "Importance_%",
            ascending=True
        )

        fig_imp = px.bar(
            importance,
            x="Importance_%",
            y="Feature",
            orientation="h",
            title=f"Model-derived Feature Importance — {selected_cqa}",
        )

        fig_imp.update_layout(
            margin=dict(
                l=40,
                r=40,
                t=60,
                b=40
            )
        )

        st.plotly_chart(
            fig_imp,
            use_container_width=True
        )

    st.markdown(
        """
        <div class="info-box">
        <b>Important:</b> Feature importance represents the contribution
        of input variables to the model's predictive behavior. It does
        not by itself prove a causal relationship between a formulation
        variable and a CQA.
        </div>
        """,
        unsafe_allow_html=True
    )